"""M2 unit tests: extraction and chunking."""

from __future__ import annotations

from threadline.ingestion.extractors import (
    detect_media_type,
    extract_document,
    chunk_passages,
    sha256_bytes,
)


def test_detect_media_type():
    assert detect_media_type("a.PDF") == "application/pdf"
    assert detect_media_type("x.unknown") is None


def test_extract_markdown_sections():
    data = b"# Root\n\nalpha body\n\n## Child\n\nbeta body"
    extracted = extract_document("doc.md", data)
    chunks = chunk_passages(extracted.passages)
    assert len(chunks) >= 1
    assert any(c.heading in ("Root", "Child") for c in chunks)


def test_extract_docx_paragraphs():
    from docx import Document
    import io
    doc = Document()
    doc.add_heading("Heading", level=1)
    doc.add_paragraph("Hello world paragraph")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    extracted = extract_document("doc.docx", data)
    chunks = chunk_passages(extracted.passages)
    assert any("Hello world" in c.content for c in chunks)
    assert any(c.heading == "Heading" for c in chunks)


def test_chunk_size_bounded():
    data = b"para\n\n" + (b"x" * 3000)
    extracted = extract_document("d.txt", data)
    chunks = chunk_passages(extracted.passages)
    assert all(len(c.content) <= 1200 for c in chunks)
    assert len(chunks) >= 2


def test_sha256_bytes_stable():
    assert sha256_bytes(b"abc") == sha256_bytes(b"abc")
    assert sha256_bytes(b"abc") != sha256_bytes(b"abd")
